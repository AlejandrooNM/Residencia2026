"""Convierte bloques de contenido en un documento Word con el formato del informe."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

CARPETA_DOCUMENTOS = Path(__file__).resolve().parents[1]
CARPETAS_IMAGENES = (
    CARPETA_DOCUMENTOS / "figuras_informe",
    CARPETA_DOCUMENTOS / "resultados",
)
FUENTE = "Times New Roman"
ESTILO_FIGURA = "Titulo Figura"
ESTILO_TABLA = "Titulo Tabla"


def _fijar_fuente(estilo, tamano, negrita=False, cursiva=False, color=(0, 0, 0)):
    estilo.font.name = FUENTE
    estilo.font.size = Pt(tamano)
    estilo.font.bold = negrita
    estilo.font.italic = cursiva
    estilo.font.color.rgb = RGBColor(*color)
    rpr = estilo.element.get_or_add_rPr()
    fuentes = rpr.get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(atributo), FUENTE)


def _campo(parrafo, instruccion: str, texto_previo: str = "") -> None:
    """Inserta un campo de Word (PAGE, TOC...) que se actualiza al abrir el documento."""
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    inicio.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instr, separador):
        parrafo.add_run()._r.append(elemento)
    if texto_previo:
        parrafo.add_run(texto_previo)
    parrafo.add_run()._r.append(fin)


class Informe:
    def __init__(self) -> None:
        self.doc = Document()
        self.figuras = 0
        self.tablas = 0
        self._configurar()

    def _configurar(self) -> None:
        seccion = self.doc.sections[0]
        seccion.page_width = Cm(21.0)
        seccion.page_height = Cm(29.7)
        seccion.top_margin = Cm(2.5)
        seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = Cm(2.5)
        seccion.right_margin = Cm(2.5)
        seccion.different_first_page_header_footer = True

        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _campo(pie, "PAGE", "1")
        for run in pie.runs:
            run.font.name = FUENTE
            run.font.size = Pt(10)

        estilos = self.doc.styles
        normal = estilos["Normal"]
        _fijar_fuente(normal, 12)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.space_after = Pt(6)

        for nombre, tamano, antes, despues in (
            ("Heading 1", 14, 0, 6),
            ("Heading 2", 14, 12, 6),
            ("Heading 3", 12, 10, 4),
        ):
            estilo = estilos[nombre]
            _fijar_fuente(estilo, tamano, negrita=True)
            formato = estilo.paragraph_format
            formato.space_before = Pt(antes)
            formato.space_after = Pt(despues)
            formato.keep_with_next = True
            formato.alignment = WD_ALIGN_PARAGRAPH.LEFT
            formato.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        for nombre in (ESTILO_FIGURA, ESTILO_TABLA):
            estilo = estilos.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
            estilo.base_style = normal
            _fijar_fuente(estilo, 11, cursiva=True)
            estilo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            estilo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            estilo.paragraph_format.space_after = Pt(12)
        estilos[ESTILO_TABLA].paragraph_format.space_before = Pt(10)
        estilos[ESTILO_TABLA].paragraph_format.space_after = Pt(4)
        estilos[ESTILO_TABLA].paragraph_format.keep_with_next = True

        for nivel in (1, 2, 3):
            try:
                estilo = estilos[f"TOC {nivel}"]
            except KeyError:
                estilo = estilos.add_style(f"TOC {nivel}", WD_STYLE_TYPE.PARAGRAPH)
            _fijar_fuente(estilo, 12, negrita=(nivel == 1))
            estilo.paragraph_format.left_indent = Cm(0.6 * (nivel - 1))
            estilo.paragraph_format.space_after = Pt(2)
            estilo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        configuracion = self.doc.settings.element
        actualizar = OxmlElement("w:updateFields")
        actualizar.set(qn("w:val"), "true")
        configuracion.append(actualizar)

    # ------------------------------------------------------------------ texto
    def _runs(self, parrafo, texto: str, tamano: float = 12, cursiva: bool = False, fuente: str = FUENTE) -> None:
        """Agrega texto; los fragmentos entre ** van en negritas."""
        for indice, trozo in enumerate(texto.split("**")):
            if not trozo:
                continue
            run = parrafo.add_run(trozo)
            run.bold = indice % 2 == 1
            run.italic = cursiva
            run.font.name = fuente
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), fuente)
            run.font.size = Pt(tamano)

    def _leyenda(self, parrafo, etiqueta: str, texto: str) -> None:
        """Leyenda con la etiqueta en negritas y el texto en cursiva, como en el borrador."""
        run = parrafo.add_run(f"{etiqueta} ")
        run.bold = True
        run.italic = False
        run.font.name = FUENTE
        run.font.size = Pt(11)
        self._runs(parrafo, texto, tamano=11, cursiva=True)

    def parrafo(self, texto, sangria=False, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, tamano=12,
                cursiva=False, despues=6):
        p = self.doc.add_paragraph()
        p.alignment = alineacion
        p.paragraph_format.space_after = Pt(despues)
        if sangria:
            p.paragraph_format.first_line_indent = Cm(1.25)
        self._runs(p, texto, tamano, cursiva)
        return p

    def centrado(self, texto, tamano=12, negrita=False, despues=4, alineacion=WD_ALIGN_PARAGRAPH.CENTER,
                 fuente=FUENTE):
        p = self.doc.add_paragraph()
        p.alignment = alineacion
        p.paragraph_format.space_after = Pt(despues)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        self._runs(p, f"**{texto}**" if negrita and texto else texto, tamano, fuente=fuente)
        return p

    def extremos(self, izquierda, derecha, tamano=12, fuente="Arial", despues=4):
        """Una línea con texto pegado al margen izquierdo y otro al derecho."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(despues)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        ancho = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
        p.paragraph_format.tab_stops.add_tab_stop(ancho, alignment=2)
        self._runs(p, f"{izquierda}\t{derecha}", tamano, fuente=fuente)
        return p

    def imagen_centrada(self, nombre, ancho_cm, despues=6):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(despues)
        p.add_run().add_picture(str(self._ruta_imagen(nombre)), width=Cm(ancho_cm))
        return p

    def vineta(self, texto, nivel=0):
        p = self.doc.add_paragraph(style="List Bullet" if nivel == 0 else "List Bullet 2")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)
        self._runs(p, texto)
        return p

    def numerada(self, numero, texto):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        self._runs(p, f"{numero}. {texto}")
        return p

    def nota(self, texto):
        p = self.parrafo(texto, sangria=False, tamano=11, cursiva=True)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        return p

    def definicion(self, termino, texto):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        self._runs(p, f"**{termino}.** {texto}")
        return p

    def codigo(self, texto):
        for linea in texto.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(linea if linea else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def titulo(self, texto, nivel, nueva_pagina=None):
        if nueva_pagina is None:
            nueva_pagina = nivel == 1
        if nueva_pagina and len(self.doc.paragraphs) > 1:
            self.salto()
        return self.doc.add_heading(texto, level=nivel)

    def titulo_sin_indice(self, texto):
        """Título con aspecto de capítulo que no aparece en el índice (páginas de índices)."""
        if len(self.doc.paragraphs) > 1:
            self.salto()
        return self.centrado(texto, tamano=14, negrita=True, despues=12, alineacion=WD_ALIGN_PARAGRAPH.LEFT)

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------- figuras y tablas
    def _ruta_imagen(self, nombre: str) -> Path | None:
        ruta = Path(nombre)
        if ruta.is_absolute():
            return ruta if ruta.exists() else None
        for carpeta in CARPETAS_IMAGENES:
            candidata = carpeta / nombre
            if candidata.exists():
                return candidata
        return None

    def figura(self, nombre, leyenda, ancho=6.0):
        self.figuras += 1
        ruta = self._ruta_imagen(nombre)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if ruta is None:
            raise FileNotFoundError(f"No se encontró la imagen {nombre}")
        p.add_run().add_picture(str(ruta), width=Inches(min(ancho, 6.2)))
        pie = self.doc.add_paragraph(style=ESTILO_FIGURA)
        self._leyenda(pie, f"Figura {self.figuras}.", leyenda)
        return self.figuras

    def figura_pendiente(self, leyenda, indicacion):
        """Recuadro vacío para imágenes institucionales que se insertan a mano."""
        self.figuras += 1
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = t.rows[0].cells[0]
        celda.width = Cm(14)
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for _ in range(5):
            p.add_run().add_break()
        self._runs(p, f"[{indicacion}]", tamano=11, cursiva=True)
        for _ in range(5):
            p.add_run().add_break()
        pie = self.doc.add_paragraph(style=ESTILO_FIGURA)
        pie.paragraph_format.space_before = Pt(4)
        self._leyenda(pie, f"Figura {self.figuras}.", leyenda)
        return self.figuras

    def tabla(self, leyenda, encabezados, filas, anchos=None, tamano=10):
        self.tablas += 1
        cabeza = self.doc.add_paragraph(style=ESTILO_TABLA)
        self._leyenda(cabeza, f"Tabla {self.tablas}.", leyenda)

        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, texto in enumerate(encabezados):
            celda = t.rows[0].cells[i]
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(2)
            self._runs(p, f"**{texto}**", tamano)
            sombra = OxmlElement("w:shd")
            sombra.set(qn("w:val"), "clear")
            sombra.set(qn("w:fill"), "D9E2F3")
            celda._tc.get_or_add_tcPr().append(sombra)
        fila_xml = t.rows[0]._tr
        propiedades = fila_xml.get_or_add_trPr()
        repetir = OxmlElement("w:tblHeader")
        repetir.set(qn("w:val"), "true")
        propiedades.append(repetir)

        for fila in filas:
            celdas = t.add_row().cells
            for i, valor in enumerate(fila):
                p = celdas[i].paragraphs[0]
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._runs(p, str(valor), tamano)
        if anchos:
            for fila in t.rows:
                for i, ancho in enumerate(anchos):
                    fila.cells[i].width = Cm(ancho)
        espacio = self.doc.add_paragraph()
        espacio.paragraph_format.space_after = Pt(6)
        espacio.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        return self.tablas

    # ---------------------------------------------------------------- índices
    def indice(self, instruccion, texto="Actualice el campo (clic derecho → Actualizar campo)."):
        p = self.doc.add_paragraph()
        _campo(p, instruccion, texto)
        return p

    def indice_contenido(self):
        return self.indice('TOC \\o "1-3" \\h \\z \\u')

    def indice_figuras(self):
        return self.indice(f'TOC \\h \\z \\t "{ESTILO_FIGURA},1"')

    def indice_tablas(self):
        return self.indice(f'TOC \\h \\z \\t "{ESTILO_TABLA},1"')

    # --------------------------------------------------------------- bloques
    def agregar(self, bloques) -> None:
        for bloque in bloques:
            tipo, *datos = bloque
            if tipo == "h1":
                self.titulo(datos[0], 1)
            elif tipo == "h1n":
                self.titulo(datos[0], 1, nueva_pagina=False)
            elif tipo == "h2":
                self.titulo(datos[0], 2)
            elif tipo == "h3":
                self.titulo(datos[0], 3)
            elif tipo == "sub":
                p = self.parrafo(f"**{datos[0]}**", despues=2)
                p.paragraph_format.keep_with_next = True
            elif tipo == "p":
                self.parrafo(datos[0])
            elif tipo == "ps":
                self.parrafo(datos[0], sangria=False)
            elif tipo == "li":
                self.vineta(datos[0])
            elif tipo == "li2":
                self.vineta(datos[0], nivel=1)
            elif tipo == "num":
                self.numerada(datos[0], datos[1])
            elif tipo == "nota":
                self.nota(datos[0])
            elif tipo == "def":
                self.definicion(datos[0], datos[1])
            elif tipo == "codigo":
                self.codigo(datos[0])
            elif tipo == "fig":
                self.figura(*datos)
            elif tipo == "figph":
                self.figura_pendiente(*datos)
            elif tipo == "tabla":
                self.tabla(*datos)
            elif tipo == "salto":
                self.salto()
            else:
                raise ValueError(f"Bloque desconocido: {tipo}")

    def guardar(self, ruta: Path) -> None:
        self.doc.save(str(ruta))
